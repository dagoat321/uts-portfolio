import sys
from docx import Document

doc = Document()
doc.add_heading('Part III: Essay (Introduction)', level=1)
doc.add_paragraph('Every person holds narratives, thoughts, and truths that have yet to be discovered. This final requirement serves as a medium for that discovery, and "Unfolding Me" represents the process of introspection.')
doc.add_paragraph('Taking the "Understanding the Self" course this semester served as a profound reflective experience. Each module revealed different facets of my identity, ranging from how Eastern and Western cultures have influenced my core values to how my physical form, beliefs, possessions, and social relationships construct my overall identity.')
doc.add_paragraph('The title Unfolding Me encapsulates the gradual, and at times challenging, process of self-examination. Similar to unfolding origami, previous experiences leave lasting marks that serve as evidence of the journey that shaped my present self.')
doc.add_paragraph('This portfolio is a compilation of these reflective insights: six post-assessment modules, a SMART goal reflection, and key realizations to carry forward. It presents an honest and structured evaluation of my personal development.')

doc.add_heading('Part IV: SMART GOAL Reflection', level=1)
doc.add_paragraph('At midterm, I established a goal to improve my time management by completing all academic requirements at least two days before their respective deadlines. This aimed to reduce last-minute cramming and foster an environment suitable for genuine reflection rather than rushed submissions.')
doc.add_paragraph('The goal was achieved partially, which provided a valuable learning experience. While there were weeks of early submissions, procrastination occasionally resurfaced. The primary insight gained is that breaking larger tasks into smaller, manageable steps yields the best results. The objective was appropriate, but the execution system required further refinement.')
doc.add_paragraph('Moving forward, pairing the deadline objective with a daily micro-habit, such as dedicating 20 minutes each morning to pending coursework, will be beneficial. Building consistency is essential. Planning ahead ensures a more structured and stress-free academic experience.')

doc.add_heading('Part XI: Essay/Parting Words', level=1)
doc.add_paragraph('Reflecting on the beginning of this semester, the course "Understanding the Self" has genuinely transformed my perspective on my identity in substantial and meaningful ways.')
doc.add_paragraph('I have learned that the self is not a static entity. It is a dynamic construction influenced by culture, upbringing, relationships, beliefs, possessions, and physical presence. I embody a blend of Eastern and Western values, philosophical and psychological insights, social and individual traits, as well as physical and spiritual dimensions.')
doc.add_paragraph('The most significant realization I take away is that self-awareness is a continuous practice rather than a final destination. This semester provided the analytical framework necessary to continually assess my identity and future aspirations.')
doc.add_paragraph('Moving forward, I am committed to continuous personal development. I resolve to strive for a version of myself that is inquisitive, compassionate, and intentional. I will prioritize effective time management, lead with empathy, and recognize that personal growth requires patience and sustained effort.')
doc.add_paragraph('I am profoundly grateful for the insights gained from this course and the unexpected depth of self-discovery it facilitated.')

doc.save('Plagscan_Submission_Parts_III_IV_XI.docx')
print('Docx created successfully.')
